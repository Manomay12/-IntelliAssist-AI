"""
Sample documents generator for IntelliAssist AI.
Creates realistic academic & technical PDF, DOCX, and TXT files for 1-click evaluation.
"""

import io
from pathlib import Path
import docx
from utils.config import SAMPLE_DATA_DIR

def create_sample_docx(filepath: Path):
    """Create a structured sample DOCX project report."""
    doc = docx.Document()
    doc.add_heading("Major Project Report: IntelliAssist AI Document Intelligence", level=0)
    
    doc.add_heading("1. Executive Summary", level=1)
    doc.add_paragraph(
        "IntelliAssist AI represents a next-generation Document Intelligence and Conversational Retrieval-Augmented "
        "Generation (RAG) platform. Modern students, enterprise analysts, and researchers face severe information overload "
        "when querying large document repositories. Traditional keyword search fails to capture semantic meaning and intent. "
        "Our system combines multi-format parsing, recursive text chunking, dense vector embeddings, cosine vector indexing, "
        "and large language model synthesis to deliver verifiable answers with exact source citations."
    )
    
    doc.add_heading("2. Core System Architecture", level=1)
    doc.add_paragraph(
        "The architecture is structured into four distinct modular pipelines: "
        "1. Ingestion Pipeline: Supports PDF, DOCX, and TXT extraction with layout normalization. "
        "2. Chunking & Indexing Pipeline: Partitions documents into 500-character overlapping chunks while preserving page metadata. "
        "3. Semantic Vector Store: Embeds chunk tokens into 384-dimensional vector space using cosine distance indexing. "
        "4. RAG Synthesis Pipeline: Formulates dynamic context prompts and prompts LLMs for structured answers with page citations."
    )
    
    doc.add_heading("3. Performance Benchmarks & Results", level=1)
    table = doc.add_table(rows=1, cols=4)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Component"
    hdr_cells[1].text = "Metric"
    hdr_cells[2].text = "Baseline"
    hdr_cells[3].text = "IntelliAssist AI"
    
    data = [
        ("Retrieval Latency", "Search Time (ms)", "145 ms", "18 ms"),
        ("Answer Faithfulness", "F1 Score (%)", "71.2%", "94.6%"),
        ("Source Attribution", "Accuracy (%)", "65.0%", "98.2%"),
        ("Vector DB Throughput", "Chunks/sec", "250", "1,850")
    ]
    for row_data in data:
        row_cells = table.add_row().cells
        for i, val in enumerate(row_data):
            row_cells[i].text = val
            
    doc.add_heading("4. Future Enhancements & Conclusion", level=1)
    doc.add_paragraph(
        "In conclusion, IntelliAssist AI provides a robust, scalable, and highly accurate solution for automated "
        "document understanding. Future work will integrate multi-modal chart comprehension, tabular question answering, "
        "and distributed vector storage clustering."
    )
    
    doc.save(filepath)

def create_sample_txt(filepath: Path):
    """Create a detailed sample text file on AI in Healthcare."""
    content = """================================================================================
AI IN HEALTHCARE & CLINICAL DECISION SUPPORT: A COMPREHENSIVE OVERVIEW
================================================================================

1. INTRODUCTION & PROBLEM STATEMENT
Artificial Intelligence (AI) and Machine Learning (ML) algorithms are transforming clinical diagnostics,
personalized medicine, and hospital administrative operations. Clinicians are overwhelmed with hundreds of
pages of Electronic Health Records (EHRs), patient clinical notes, and genomic sequencing data. 
Automated NLP systems assist doctors by summarizing patient history and highlighting diagnostic risk factors.

2. CLINICAL APPLICATIONS
A. Medical Imaging: Deep Convolutional Neural Networks (CNNs) analyze radiographs, MRI scans, and CT scans
   with diagnostic accuracy exceeding 92.5%, particularly in early oncology screening and thoracic pathology.
B. Clinical NLP: Transformer models extract structured medical codes (ICD-10, SNOMED-CT) from unstructured
   physician consultation notes, decreasing documentation overhead by 45%.
C. Predictive Risk Stratification: Recurrent neural networks and gradient-boosted trees predict sepsis onset
   up to 6 hours in advance, significantly lowering ICU mortality rates.

3. REGULATORY COMPLIANCE & ETHICAL SAFETY
Clinical AI systems must adhere to strict regulatory standards including HIPAA, GDPR, and FDA Software as a
Medical Device (SaMD) guidelines. Model explainability, source verification, and bias mitigation remain 
critical requirements before deploying diagnostic AI at scale.

4. CONCLUSION
Integrating semantic search and RAG architecture enables medical researchers to query thousands of biomedical
journals rapidly, expediting clinical discovery and improving patient outcomes worldwide.
"""
    with open(filepath, "w", encoding="utf-8") as f:
        f.write(content)

def create_sample_pdf(filepath: Path):
    """Generate a clean sample PDF document about Transformers and RAG using a minimal valid PDF generator."""
    # Build standard PDF format bytes
    lines_p1 = [
        "RESEARCH PAPER: TRANSFORMERS & RETRIEVAL-AUGMENTED GENERATION",
        "",
        "Abstract:",
        "Large Language Models (LLMs) often suffer from factual hallucinations and knowledge cutoff limitations.",
        "Retrieval-Augmented Generation (RAG) mitigates these challenges by grounding generative responses in",
        "authoritative external document repositories. This paper reviews self-attention mechanisms, vector indexing,",
        "and top-k semantic context ranking.",
        "",
        "1. Attention Mechanism in Transformer Architectures:",
        "The scaled dot-product attention function maps queries, keys, and values to an output vector weighted by",
        "softmax compatibility: Attention(Q, K, V) = softmax(Q * K^T / sqrt(d_k)) * V.",
        "Multi-Head Attention enables the model to jointly attend to information from different representation subspaces.",
        "",
        "2. Vector Indexing and Semantic Similarity:",
        "Embeddings map high-dimensional text chunks into dense vector representations. Cosine similarity calculates",
        "the directional alignment between query embeddings and document chunk embeddings, enabling rapid nearest-neighbor retrieval."
    ]
    
    lines_p2 = [
        "3. Retrieval-Augmented Generation (RAG) Pipeline:",
        "RAG combines parametric memory (pre-trained LLM weights) with non-parametric memory (vector database).",
        "When a user submits a query, the retriever fetches the most relevant text chunks from the vector store.",
        "The context is appended to the system prompt, instructing the LLM to formulate an evidence-backed answer.",
        "",
        "4. Experimental Evaluation & Results:",
        "Our experiments demonstrate that RAG improves response factual accuracy by 38.4% and reduces hallucination",
        "rates from 24.1% to less than 2.3% across open-domain technical benchmarks.",
        "",
        "5. Conclusion:",
        "RAG represents a foundational architecture for enterprise AI and academic document assistants.",
        "Source citations provide transparency and verifiable accountability for generative AI outputs."
    ]

    def build_pdf_stream(pages_text_list):
        objects = []
        # obj 1: Catalog
        # obj 2: Pages
        # For each page: Page obj, Content stream obj
        page_objs = []
        
        # We will dynamically construct objects
        # 1: Catalog
        # 2: Pages
        # 3: Font
        font_obj = "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"
        
        content_objs = []
        for lines in pages_text_list:
            stream_body = "BT\n/F1 11 Tf\n50 750 Td\n15 TL\n"
            for line in lines:
                # Escape parentheses
                safe_line = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
                stream_body += f"({safe_line}) '\n"
            stream_body += "ET"
            content_objs.append(stream_body)

        total_pages = len(pages_text_list)
        # Obj numbers:
        # 1 = Catalog
        # 2 = Pages
        # 3 = Font
        # For i in range(total_pages):
        # 4 + 2*i = Page i
        # 5 + 2*i = Content i
        
        pdf = "%PDF-1.4\n"
        offsets = {}
        
        def add_obj(num, content):
            nonlocal pdf
            offsets[num] = len(pdf.encode("latin1"))
            pdf += f"{num} 0 obj\n{content}\nendobj\n"

        add_obj(1, "<< /Type /Catalog /Pages 2 0 R >>")
        
        kids_refs = " ".join([f"{4 + 2*i} 0 R" for i in range(total_pages)])
        add_obj(2, f"<< /Type /Pages /Kids [{kids_refs}] /Count {total_pages} >>")
        add_obj(3, font_obj)

        for i in range(total_pages):
            page_num = 4 + 2*i
            content_num = 5 + 2*i
            c_body = content_objs[i]
            c_len = len(c_body.encode("latin1"))
            
            add_obj(page_num, f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_num} 0 R /Resources << /Font << /F1 3 0 R >> >> >>")
            add_obj(content_num, f"<< /Length {c_len} >>\nstream\n{c_body}\nendstream")

        xref_pos = len(pdf.encode("latin1"))
        total_objs = 3 + 2 * total_pages
        pdf += f"xref\n0 {total_objs + 1}\n0000000000 65535 f \n"
        for i in range(1, total_objs + 1):
            pdf += f"{offsets[i]:010d} 00000 n \n"
            
        pdf += f"trailer\n<< /Size {total_objs + 1} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF"
        return pdf.encode("latin1")

    pdf_bytes = build_pdf_stream([lines_p1, lines_p2])
    with open(filepath, "wb") as f:
        f.write(pdf_bytes)

def generate_all_samples():
    """Ensure all sample files exist in sample_data directory."""
    SAMPLE_DATA_DIR.mkdir(parents=True, exist_ok=True)
    pdf_path = SAMPLE_DATA_DIR / "AI_Research_Paper_Transformers_and_RAG.pdf"
    docx_path = SAMPLE_DATA_DIR / "Project_Report_IntelliAssist_AI.docx"
    txt_path = SAMPLE_DATA_DIR / "Machine_Learning_Healthcare_Overview.txt"

    if not pdf_path.exists():
        create_sample_pdf(pdf_path)
    if not docx_path.exists():
        create_sample_docx(docx_path)
    if not txt_path.exists():
        create_sample_txt(txt_path)

    return [pdf_path, docx_path, txt_path]
